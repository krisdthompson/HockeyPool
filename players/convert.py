"""Convert the organizer's pick sheet (.xlsx) into players.csv for the app.

    pip install openpyxl
    python3 players/convert.py "players/<pick sheet>.xlsx" > players/players.csv

If players/espn.csv exists (see espn_parse.py), ESPN's tiers become an
"Expert" projection: ESPN's Nth forward gets the Nth-highest forward points
total from last season (defensemen likewise). ESPN's lists are added to the
Note column.

Stats are last season's. Injury "out until" dates become expected games
missed, assuming the season opens SEASON_START and runs 82 games over
SEASON_DAYS days. ROOKIES and PROJ_OVERRIDES are hand-maintained guesses;
fix them here or in the app.
"""
import csv
import datetime as dt
import math
import os
import re
import sys
import unicodedata

import openpyxl

SEASON_START = dt.date(2026, 10, 7)
SEASON_DAYS = 186

# Under 25 NHL games last season and (to my knowledge) still rookie-eligible.
ROOKIES = {
    "Ryan Ufko", "Porter Martone", "Cole Hutson", "Anton Frondell", "Alex Bump",
    "Konsta Helenius", "Ilya Protas", "Samuel Honzek", "Bradly Nadeau",
    "Carter Yakemchuk", "James Hagens", "Victor Eklund", "Michael Brandsegg-Nygard",
    "Tristan Luneau", "Luca Cagnoni", "Gavin McKenna", "Ivar Stenberg",
    "Roman Kantserov", "Viggo Bjorck", "Tij Iginla", "Vitali Pinchuk",
    "Trevor Connelly", "Charlie Stramel",
}

# Players whose last-season line says nothing useful about this season.
PROJ_OVERRIDES = {
    "Aleksander Barkov": (70, "Missed last season (knee); ~80-pt pace before that"),
}


ESPN_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "espn.csv")
ESPN_TEAMS = {"SJ": "SJS", "TB": "TBL", "MON": "MTL", "LA": "LAK", "NJ": "NJD", "WAS": "WSH", "UTAH": "UTA"}
# ESPN spelling -> organizer spelling, where first names differ.
ESPN_ALIASES = {"aliaksei protas": "alexei protas", "matt boldy": "matthew boldy",
                "matty beniers": "matthew beniers", "braden schneider": "brayden schneider",
                "jj moser": "janis moser"}
ESPN_LABELS = {"Flag": "ESPN must-draft", "Sleeper": "ESPN sleeper", "Bounceback": "ESPN bounceback",
               "Breakout": "ESPN breakout", "Rookie": "ESPN rookie to know",
               "LateRound": "ESPN late-round value", "PointsLeague": "ESPN: better in points leagues"}


def name_key(n: str) -> str:
    n = unicodedata.normalize("NFKD", n).encode("ascii", "ignore").decode().lower().replace(".", "")
    n = " ".join(n.split())
    return ESPN_ALIASES.get(n, n)


def load_espn(players):
    """Returns {name_key: (expert_pts, [note parts])} for players on ESPN's sheet."""
    if not os.path.exists(ESPN_FILE):
        return {}
    fwd = lambda p: p["pos"] != "D"
    pts = {grp: sorted((p["pts"] or 0 for p in players if fwd(p) == (grp == "F")), reverse=True) for grp in ("F", "D")}
    out = {}
    for e in csv.DictReader(open(ESPN_FILE)):
        k = name_key(e["Name"])
        expert, notes = out.get(k, (None, []))
        if e["List"] in ("F", "D"):
            rank = int(e["Order"])
            expert = pts[e["List"]][rank - 1]
            notes = [f"ESPN {e['List']} tier {e['Tier']} (#{rank})"] + notes
        elif e["List"] in ESPN_LABELS:
            notes = notes + [ESPN_LABELS[e["List"]]]
        out[k] = (expert, notes)
    return out


def games_missed(until: dt.date) -> int:
    days = (until - SEASON_START).days
    return max(0, math.ceil(days * 82 / SEASON_DAYS))


def parse_until(status: str):
    m = re.search(r"until at least (\w{3}) (\d{1,2})", status)
    if not m:
        return None
    month = dt.datetime.strptime(m.group(1), "%b").month
    year = 2026 if month >= 7 else 2027
    return dt.date(year, month, int(m.group(2)))


def injury_fields(kind: str, status: str):
    """Returns (injury label, games missed, note)."""
    kind, status = (kind or "").strip(), (status or "").strip()
    if not kind and not status:
        return "", 0, ""
    if kind == "Suspension":
        m = re.match(r"(\d+)", status)
        n = int(m.group(1)) if m else 0
        return f"Susp {n}g", n, status
    if status.startswith("Day-to-Day"):
        return f"DTD {kind}", 0, ""
    until = parse_until(status)
    if until is None:
        return kind, 0, status
    miss = games_missed(until)
    when = until.strftime("%b %-d")
    if miss == 0:
        # Back before opening night: keep it visible, day-to-day level risk.
        return f"Camp: {kind}, back ~{when}", 0, ""
    return f"Out: {kind}, to ~{when}", miss, ""


def main(path: str):
    ws = openpyxl.load_workbook(path, data_only=True).active
    header = [str(c.value or "").strip() for c in ws[1]]
    col = {h: i for i, h in enumerate(header) if h}
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[col["PLAYER NAME"]] and r[col["POS"]]]
    espn = load_espn([{"pos": r[col["POS"]], "pts": r[col["Pts"]]} for r in rows])
    matched = set()
    out = csv.writer(sys.stdout, lineterminator="\n")
    out.writerow(["Name", "Team", "Pos", "GP", "G", "A", "Pts", "Proj", "Expert", "Miss", "Rookie", "Injury", "Note"])
    for row in rows:
        name = row[col["PLAYER NAME"]]
        if not name or not row[col["POS"]]:
            continue  # skips the Spent/Left rows
        name = str(name).strip()
        team = row[col["TEAM"]]
        injury, miss, note = injury_fields(row[col["Injury"]], row[col["Status"]])
        proj = ""
        if name in PROJ_OVERRIDES:
            proj, note = PROJ_OVERRIDES[name]
        if team == "UFA":
            note = "Unsigned free agent"
        expert, espn_notes = espn.get(name_key(name), (None, []))
        if name_key(name) in espn:
            matched.add(name_key(name))
        rookie = name in ROOKIES or "ESPN rookie to know" in espn_notes
        note = "; ".join(x for x in [note] + espn_notes if x)
        stat = lambda k: "" if row[col[k]] is None else row[col[k]]
        out.writerow([name, team, row[col["POS"]], stat("GP"), stat("G"), stat("A"), stat("Pts"),
                      proj, expert or "", miss or "", "Y" if rookie else "", injury, note])
    for k in sorted(set(espn) - matched):
        print(f"not on the organizer's list: {k}", file=sys.stderr)


if __name__ == "__main__":
    main(sys.argv[1])
